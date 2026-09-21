#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Bioindustrial-Park: BioSTEAM's Premier Biorefinery Models and Results
# Copyright (C) 2021-, Sarang Bhagwat <sarangbhagwat.developer@gmail.com>
#
# This module is under the UIUC open-source license. See
# github.com/BioSTEAMDevelopmentGroup/biosteam/blob/master/LICENSE.txt
# for license details.
"""Offline pure-logic test of build_opt_IRR_split12d_workbook (no load, no
simulation). Loads the generator by file path -- its biorefineries imports
live inside main(), so the import is build-free -- and checks the merge,
the workbook writer (on a synthetic template in a temp dir), the
reload-verify comparison and the ScenarioSpec block formatter.
Exit 0 + ALL CHECKS PASSED = clean."""
import os
import sys
import math
import tempfile
import importlib.util

from openpyxl import Workbook, load_workbook

PKG_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _load(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


gen = _load('build_opt_IRR_split12d_workbook',
            os.path.join(PKG_DIR, 'analyses',
                         'build_opt_IRR_split12d_workbook.py'))
checks = []


def check(name, cond):
    checks.append((name, bool(cond)))
    print(f'  [{"PASS" if cond else "FAIL"}] {name}')


def raises(exc_type, func, *args, **kwargs):
    try:
        func(*args, **kwargs)
    except exc_type:
        return True
    return False


# --- import is build-free ---------------------------------------------------
check('importing the generator does not import biorefineries',
      not any(m == 'biorefineries' or m.startswith('biorefineries.')
              for m in sys.modules))
check('settings default to the opt_IRR / trial-1602 relocation',
      (gen.SCENARIO_NAME, gen.OBJECTIVE_NAME, gen.ANCHOR_SCENARIO,
       gen.TRIAL_NUMBER) == ('opt_IRR', 'IRR', 'A', 1602)
      and 'metabolic_split_12d' in gen.STUDY_NAME)

# --- merge_applied ----------------------------------------------------------
a_snapshot = {'k_3': 5.81, 'k_7': 1.203, 'k_17': 0.1077, 'K_17': 0.0086}
applied = {'k_3': 1.278, 'k_17': 2.154, 'threshold_conc': 170.9,
           'target_delta': 5.0, 'max_n_spikes': 50}
merged = gen.merge_applied(a_snapshot, applied)
check('merge keeps exactly the snapshot keys (feeding keys dropped)',
      set(merged) == set(a_snapshot))
check('merge: sampled names take the trial value',
      merged['k_3'] == 1.278 and merged['k_17'] == 2.154)
check('merge: non-sampled names (k_7, K_*) keep the snapshot value',
      merged['k_7'] == 1.203 and merged['K_17'] == 0.0086)
check('merge: every value is a Python float',
      all(type(v) is float for v in merged.values()))
check('merge: an unknown non-feeding name raises ValueError',
      raises(ValueError, gen.merge_applied, a_snapshot, {'k_99': 1.0}))

# --- write_workbook on a synthetic template ---------------------------------
HEADER = ['Parameter name', 'Element', 'Kind', 'Units', 'Baseline', 'Shape',
          'Lower', 'Midpoint', 'Upper', 'References', 'Load statement']
ROWS = [
    ['Plant annual operating days', 'TEA', 'isolated', 'd', 330, 'Triangular',
     297, 330, 363, None, 'tea.operating_days = x'],
    ['Isobutanol selling unit price (sale)', 'TEA', 'isolated', '$/kg', 1.5,
     'Triangular', 0.9, 1.5, 1.8, 'price note', 'V514.isobutanol_price = x'],
    ['._k_3', 'Kinetics', 'coupled', 'g_per_l_per_h', 5.81, 'triangular',
     4.648, 5.81, 6.972, None, 'V406.nsk_kinetic_model._te.k_3 = x'],
    ['._k_17', 'Kinetics', 'coupled', 'g_per_l_per_h', 44, 'triangular',
     35.2, 44, 52.8, 'k_17 = 44 g/L/h: old note',
     'V406.nsk_kinetic_model._te.k_17 = x'],
]


def _rows(path):
    ws = load_workbook(path, data_only=True).active
    return {r[10]: r for r in ws.iter_rows(min_row=2, values_only=True)}


with tempfile.TemporaryDirectory() as td:
    template = os.path.join(td, 'template.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    wb.save(template)

    out = os.path.join(td, 'out.xlsx')
    n_kin = gen.write_workbook(template, out, merged, 2.0,
                               provenance='trial 1602')
    rows = _rows(out)
    k3 = rows['V406.nsk_kinetic_model._te.k_3 = x']
    k17 = rows['V406.nsk_kinetic_model._te.k_17 = x']
    price = rows['V514.isobutanol_price = x']
    days = rows['tea.operating_days = x']
    check('write_workbook returns the kinetic-row count', n_kin == 2)
    check('kinetic row Baseline <- full_applied value',
          k3[4] == 1.278 and k17[4] == 2.154)
    check('kinetic row is Triangular +/-20 %',
          k17[5] == 'Triangular'
          and math.isclose(k17[6], 0.8*2.154) and math.isclose(k17[7], 2.154)
          and math.isclose(k17[8], 1.2*2.154))
    check('kinetic row References is prefixed with the provenance',
          k17[9] == 'trial 1602 | template note: k_17 = 44 g/L/h: old note'
          and k3[9] == 'trial 1602')
    check('price row Baseline <- ibo_price, relative spread preserved',
          price[4] == 2.0 and math.isclose(price[6], 0.9*2.0/1.5)
          and math.isclose(price[7], 2.0) and math.isclose(price[8], 1.8*2.0/1.5)
          and price[9] == 'price note')
    check('non-kinetic rows are untouched', list(days) == ROWS[0])
    check('template itself is not modified',
          _rows(template)['V406.nsk_kinetic_model._te.k_17 = x'][4] == 44)

    out2 = os.path.join(td, 'out2.xlsx')
    incomplete = {k: v for k, v in merged.items() if k != 'k_17'}
    check('a kinetic row missing from full_applied raises ValueError',
          raises(ValueError, gen.write_workbook, template, out2, incomplete,
                 2.0))
    check('... and nothing is written', not os.path.exists(out2))

# --- assert_reload_matches --------------------------------------------------
reproduced = dict(ethanol=0.5, isobutanol=1.2, objective=0.2729)
close = dict(ethanol=0.5002, isobutanol=1.2003, objective=0.27295)
deltas = gen.assert_reload_matches(reproduced, close, 1e-3)
check('reload deltas returned per key, all within tol',
      set(deltas) == {'ethanol', 'isobutanol', 'objective'}
      and all(0.0 < d < 1e-3 for d in deltas.values()))
far = dict(close, isobutanol=1.25)
check('a delta beyond tol raises RuntimeError',
      raises(RuntimeError, gen.assert_reload_matches, reproduced, far, 1e-3))
check('a non-finite value raises RuntimeError',
      raises(RuntimeError, gen.assert_reload_matches, reproduced,
             dict(close, ethanol=math.nan), 1e-3))

# --- format_scenario_spec ---------------------------------------------------
block = gen.format_scenario_spec(
    'opt_IRR', 'parameter-distributions_corn_IBO_EtOH_opt_IRR.xlsx',
    dict(threshold=170.90680171931479, target=175.90680171931479,
         spike=600.0, max_n_spikes=50),
    dict(ethanol=0.5, isobutanol=1.2), 'IRR', 0.2729)
registry = eval('{' + block + '}', {'ScenarioSpec': dict})
entry = registry['opt_IRR']
check('ScenarioSpec block evaluates to the expected entry',
      entry == dict(
          name='opt_IRR',
          workbook='parameter-distributions_corn_IBO_EtOH_opt_IRR.xlsx',
          max_n_spikes=50, threshold_conc=170.90680171931479,
          target_conc=175.90680171931479,
          expected={'ethanol': 0.5, 'isobutanol': 1.2},
          objective_name='IRR', objective_value=0.2729))
check('block never sets spike_conc / stage_1_max_x (they stay None)',
      'spike_conc' not in block and 'stage_1_max_x' not in block)

n_fail = sum(1 for _n, ok in checks if not ok)
print()
if n_fail:
    print(f'{n_fail} CHECK(S) FAILED')
    raise SystemExit(1)
print(f'ALL {len(checks)} CHECKS PASSED')
