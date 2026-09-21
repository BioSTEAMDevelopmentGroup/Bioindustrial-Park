#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Bioindustrial-Park: BioSTEAM's Premier Biorefinery Models and Results
# Copyright (C) 2021-, Sarang Bhagwat <sarangbhagwat.developer@gmail.com>
#
# This module is under the UIUC open-source license. See
# github.com/BioSTEAMDevelopmentGroup/biosteam/blob/master/LICENSE.txt
# for license details.
"""One-shot generator (ASK-FIRST to run; it simulates) that relocates an
opt_* scenario to ONE recorded trial of an ethanol_isobutanol x
metabolic_split_12d kinetic-optimization study (spec docs/superpowers/specs/
2026-09-20-opt-IRR-relocate-to-split12d-optimum-design.md). The settings
cell defaults to opt_IRR <- trial 1602 of the 2026-09-16 PI (log-tail) GP
study (the highest-IRR point of the campaign; the scenario's regression
objective stays IRR).

It:
  1. isobutanol.load() (default both-trains build == IBO_EtOH-only at S201
     split 1.0, the studies' build).
  2. Snapshots the anchor scenario's full live kinetics + isobutanol price.
  3. Reproduces the trial with kinetic_optimization.reproduce_split12d_trial
     (mode='both', A-referenced burden ON, restore=True) and aborts unless
     the reproduction is clean (no error, no cross-check mismatch, no metric
     warning, recorded state COMPLETE).
  4. Merges the trial's applied kinetics over the anchor snapshot: k_7 / k_8,
     every K_* and the non-sampled rates keep the anchor's INTENDED
     (non-derated) values -- the active burden derates k_7 / k_8 at the
     load_simulate choke point, exactly as during the study.
  5. Backs up the existing scenario workbook (once), then writes the new one:
     a copy of the B workbook with every kinetic row's Baseline <- the merged
     value (Triangular +/-20 %), the isobutanol-price row <- the anchor's
     price (B's relative spread preserved), all other rows kept.
  6. Reloads the written workbook and re-simulates under the still-active
     burden; aborts unless the MPSPs and objective match the reproduction
     within RELOAD_VERIFY_TOL (defends against openpyxl save issues).
  7. Prints the ScenarioSpec block to paste into scenarios.SCENARIOS.

Trust a FRESH-KERNEL smoke test of the scenario, not this script's
in-process printout, as the go/no-go.

Every biorefineries import lives inside main(), so importing this module is
build-free (analyses/test_build_opt_IRR_split12d_workbook_offline.py loads
it by file path).

Run (ask first; fresh kernel, never next to another simulation):
  & "C:/Users/saran/anaconda3/envs/IBO_2026/python.exe" analyses/build_opt_IRR_split12d_workbook.py
"""
import os
import re
import math
import shutil
import argparse

#%% Settings (edit and run)
#: A study name (resolved to analyses/results/<name>_trajectory.csv) or the
#: path of a trajectory CSV.
STUDY_NAME = ('kin_opt_ethanol_isobutanol_metabolic_split_12d_pi_log-tail_gp_'
              'rb0.001-4_ib0.75-1.5_aA_burden')
#: Best COMPLETE trial of that study (recorded PI 0.8088, IRR 0.2726,
#: IBO 41.74 / EtOH 29.21 g/L, tau 28.35 h, spike cap 50 / 1 actual spike).
TRIAL_NUMBER = 1602
#: Scenario that supplies every non-sampled kinetic parameter, the basis of
#: the un-referenced groups and the isobutanol price. The split_12d studies
#: ran from scenario A; never anchor to B or an opt_* scenario.
ANCHOR_SCENARIO = 'A'
#: Registry key of the scenario being relocated (its workbook is rewritten).
SCENARIO_NAME = 'opt_IRR'
#: kinetic_optimization.OBJECTIVE_REGISTRY key the scenario pins.
OBJECTIVE_NAME = 'IRR'
#: Max relative delta (MPSPs, objective) between the reproduction and the
#: reload-verify simulation of the written workbook.
RELOAD_VERIFY_TOL = 1e-3

#%% Paths and constants
_ANALYSES_DIR = os.path.dirname(os.path.abspath(__file__))
WORKBOOK_DIR = os.path.join(_ANALYSES_DIR, 'full', 'parameter_distributions')
B_WORKBOOK = os.path.join(WORKBOOK_DIR,
                          'parameter-distributions_corn_IBO_EtOH_B.xlsx')
#: Non-kinetic decision entries reproduce_split12d_trial passes through in
#: result['applied']; the feeding strategy is read from result['feeding'].
FEEDING_KEYS = ('threshold_conc', 'target_delta', 'max_n_spikes')
_LOAD_RE = re.compile(r'_te\.(\w+)\s*=')

#%% Pure helpers (no model access; covered by the offline test)
def merge_applied(a_snapshot, applied):
    """The scenario's FULL kinetic state {name: float} over exactly
    `a_snapshot`'s names: the trial's `applied` values where sampled /
    grouped, the anchor snapshot elsewhere. FEEDING_KEYS entries of
    `applied` are dropped; any other name absent from the snapshot is a
    ValueError (a kinetic name the live model does not have)."""
    unknown = [name for name in applied
               if name not in a_snapshot and name not in FEEDING_KEYS]
    if unknown:
        raise ValueError(f'applied names {unknown} are neither kinetic '
                         'parameters of the anchor snapshot nor feeding keys '
                         f'{FEEDING_KEYS}')
    return {name: float(applied.get(name, value))
            for name, value in a_snapshot.items()}


def write_workbook(template_path, out_path, full_applied, ibo_price,
                   provenance=None):
    """Copy `template_path` (the B workbook) to `out_path` with every kinetic
    row's Baseline <- full_applied[name] and a Triangular +/-20 %
    distribution, and the isobutanol-price row <- `ibo_price` with the
    template's relative spread preserved; every other row is kept. With
    `provenance`, each kinetic row's References cell becomes `provenance`
    (+ ' | template note: <old>' when the template had one), so a copied
    note such as B's "k_17 = 44 ..." cannot pass for the new Baseline's
    source. ValueError BEFORE anything is saved if a kinetic row's name is
    not in `full_applied`. Returns the number of kinetic rows written."""
    from openpyxl import load_workbook
    wb = load_workbook(template_path, data_only=True)
    ws = wb.active
    header = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
    col_load, col_base = header['Load statement'], header['Baseline']
    col_shape, col_refs = header['Shape'], header['References']
    col_low, col_mid, col_up = (header['Lower'], header['Midpoint'],
                                header['Upper'])
    kinetic_rows, price_rows = [], []
    for r in ws.iter_rows(min_row=2):
        load_stmt = r[col_load].value
        if not load_stmt:
            continue
        m = _LOAD_RE.search(str(load_stmt))
        if m:
            kinetic_rows.append((m.group(1), r))
        elif 'isobutanol_price' in str(load_stmt):
            price_rows.append(r)
    missing = [name for name, _ in kinetic_rows if name not in full_applied]
    if missing:
        raise ValueError(f'kinetic workbook rows {missing} are not in the '
                         'applied kinetic state; nothing written')
    for name, r in kinetic_rows:
        base = float(full_applied[name])
        r[col_base].value = base
        r[col_shape].value = 'Triangular'
        r[col_low].value = 0.8*base
        r[col_mid].value = 1.0*base
        r[col_up].value = 1.2*base
        if provenance:
            old = r[col_refs].value
            r[col_refs].value = (f'{provenance} | template note: {old}'
                                 if old else provenance)
    for r in price_rows:
        template_base = r[col_base].value
        r[col_base].value = float(ibo_price)
        if template_base:                  # preserve the relative spread
            scale = float(ibo_price)/float(template_base)
            for c in (col_low, col_mid, col_up):
                if r[c].value is not None:
                    r[c].value = float(r[c].value)*scale
    wb.save(out_path)
    return len(kinetic_rows)


def assert_reload_matches(reproduced, reloaded, tol):
    """{key: relative delta} over 'ethanol', 'isobutanol' (MPSPs) and
    'objective' between the reproduction and the reload-verify simulation
    (|a - b| / max(|a|, |b|), Python floats -- flexsolve's global
    np.seterr(invalid='raise') makes numpy-scalar NaN comparisons raise).
    RuntimeError naming every key whose delta exceeds `tol` or whose value
    is non-finite."""
    deltas, bad = {}, []
    for key in ('ethanol', 'isobutanol', 'objective'):
        a, b = float(reproduced[key]), float(reloaded[key])
        if not (math.isfinite(a) and math.isfinite(b)):
            deltas[key] = math.nan
            bad.append(f'{key}: reproduced {a!r} vs reloaded {b!r} '
                       '(non-finite)')
            continue
        scale = max(abs(a), abs(b))
        deltas[key] = 0.0 if scale == 0.0 else abs(a - b)/scale
        if deltas[key] > tol:
            bad.append(f'{key}: reproduced {a!r} vs reloaded {b!r} '
                       f'(rel {deltas[key]:.3g} > {tol:g})')
    if bad:
        raise RuntimeError('the written workbook does not reproduce the '
                           'trial: ' + '; '.join(bad))
    return deltas


def format_scenario_spec(name, workbook, feeding, mpsps, objective_name,
                         objective_value):
    """The paste-ready scenarios.SCENARIOS entry. `feeding` is
    reproduce_split12d_trial's result['feeding'] (threshold / target /
    spike / max_n_spikes); the spike concentration is NOT written
    (spike_conc stays None = the load() default the preset pins)."""
    return (
        f"    '{name}': ScenarioSpec(\n"
        f"        name='{name}',\n"
        f"        workbook='{workbook}',\n"
        f"        max_n_spikes={int(feeding['max_n_spikes'])}, "
        f"threshold_conc={float(feeding['threshold'])!r}, "
        f"target_conc={float(feeding['target'])!r},\n"
        f"        expected={{'ethanol': {float(mpsps['ethanol'])!r}, "
        f"'isobutanol': {float(mpsps['isobutanol'])!r}}},\n"
        f"        objective_name={objective_name!r}, "
        f"objective_value={float(objective_value)!r}),")
