#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Bioindustrial-Park: BioSTEAM's Premier Biorefinery Models and Results
# Copyright (C) 2021-, Sarang Bhagwat <sarangbhagwat.developer@gmail.com>
#
# This module is under the UIUC open-source license. See
# github.com/BioSTEAMDevelopmentGroup/biosteam/blob/master/LICENSE.txt
# for license details.
"""One-shot generator (ASK-FIRST to run; it simulates) of ONE opt_* scenario
workbook + registry pin from ONE recorded trial of an ethanol_isobutanol x
metabolic_split_12d kinetic-optimization campaign (spec docs/superpowers/
specs/2026-09-20-opt-IRR-relocate-to-split12d-optimum-design.md; generalized
2026-10-08 from build_opt_IRR_split12d_workbook.py to every opt_* scenario).
OPT_SCENARIOS maps each scenario to the best trial (by the campaign's own
objective) of one of the eight 2026-09-23/24 campaigns: the seven seed-350
GP campaigns (_rs350) and the PI (log-tail) relay (_rl15c111dc).

It:
  1. isobutanol.load() (default both-trains build == IBO_EtOH-only at S201
     split 1.0, the campaigns' build).
  2. Snapshots the anchor scenario's full live kinetics + isobutanol price.
  3. Reproduces the trial with kinetic_optimization.reproduce_split12d_trial
     (mode='both', A-referenced burden ON, restore=True) and aborts unless
     the reproduction is clean: no error, no cross-check mismatch, recorded
     state COMPLETE, and every FERMENTATION-level tracked metric within the
     reproduction tolerance of its recorded value. SYSTEM-level (TEA)
     metric warnings are reported but tolerated: they move with the
     hensmith version (HXN synthesis), and the campaigns ran under the
     2b5b27d export, not the local clone the pins are taken on.
  4. Merges the trial's applied kinetics over the anchor snapshot: k_7 / k_8,
     every K_* and the non-sampled rates keep the anchor's INTENDED
     (non-derated) values -- the active burden derates k_7 / k_8 at the
     load_simulate choke point, exactly as during the campaign.
  5. Writes the scenario workbook: a copy of the B workbook with every
     kinetic row's Baseline <- the merged value (Triangular +/-20 %), the
     isobutanol-price row <- the anchor's price (B's relative spread
     preserved), all other rows kept. (Previous versions live in git.)
  6. Reloads the written workbook and re-simulates under the still-active
     burden; aborts unless the MPSPs and EVERY tracked metric (convergence
     diagnostics excepted) match the reproduction within RELOAD_VERIFY_TOL
     -- the scenario's baseline IS the reproduced trial.
  7. Prints the ScenarioSpec block to paste into scenarios.SCENARIOS.

Trust a FRESH-KERNEL smoke test of the scenario, not this script's
in-process printout, as the go/no-go.

Every biorefineries import lives inside main(), so importing this module is
build-free (analyses/test_build_opt_split12d_workbooks_offline.py loads it
by file path).

Run (ask first; one fresh process per scenario; several may run in
parallel only on a warm numba cache):
  & "C:/Users/saran/anaconda3/envs/IBO_2026/python.exe" analyses/build_opt_split12d_workbooks.py opt_PI_TRY_informed
"""
import os
import re
import json
import math
import argparse

#%% Settings
_STEM = 'kin_opt_ethanol_isobutanol_metabolic_split_12d_'
_TAIL = '_gp_rb0.001-4_ib0.75-1.5_aA_rs350_burden'
#: scenario -> (OBJECTIVE_REGISTRY key the scenario pins, campaign name
#: (resolved to analyses/results/<name>_trajectory.csv), trial_number of the
#: campaign's best COMPLETE trial by its own objective). The PI campaigns
#: optimized PI (log-tail), whose argmax is PI's; the scenarios pin PI.
OPT_SCENARIOS = {
    # Uninformed PI campaign: never left the ethanol basin (best from #103).
    'opt_PI_uninformed': ('PI', _STEM + 'pi_log-tail' + _TAIL, 103),
    # TRY-informed PI relay over the six process-level _rs350 campaigns:
    # co-production IBO 39.6 + EtOH 40.6 g/L, PI 0.808.
    'opt_PI_TRY_informed': (
        'PI', _STEM + 'pi_log-tail_gp_rb0.001-4_ib0.75-1.5_aA_rl15c111dc_burden',
        1912),
    'opt_IBO_titer':        ('IBO titer',        _STEM + 'ibo_titer' + _TAIL, 84),
    'opt_IBO_yield':        ('IBO yield',        _STEM + 'ibo_yield' + _TAIL, 841),
    'opt_IBO_productivity': ('IBO productivity', _STEM + 'ibo_productivity' + _TAIL, 1472),
    'opt_EtOH_titer':       ('EtOH titer',       _STEM + 'etoh_titer' + _TAIL, 1820),
    'opt_EtOH_yield':       ('EtOH yield',       _STEM + 'etoh_yield' + _TAIL, 1923),
    'opt_EtOH_productivity': ('EtOH productivity', _STEM + 'etoh_productivity' + _TAIL, 365),
}
#: Scenario that supplies every non-sampled kinetic parameter, the basis of
#: the un-referenced groups and the isobutanol price. The split_12d
#: campaigns ran from scenario A; never anchor to B or an opt_* scenario.
ANCHOR_SCENARIO = 'A'
#: Max relative delta between the reproduction and the reload-verify
#: simulation of the written workbook (MPSPs + every tracked metric).
RELOAD_VERIFY_TOL = 1e-3
#: Max relative delta of a FERMENTATION-level tracked metric vs the
#: campaign's recorded value (reproduce_split12d_trial's default).
METRIC_CHECK_TOL = 0.02

#%% Paths and constants
_ANALYSES_DIR = os.path.dirname(os.path.abspath(__file__))
WORKBOOK_DIR = os.path.join(_ANALYSES_DIR, 'full', 'parameter_distributions')
B_WORKBOOK = os.path.join(WORKBOOK_DIR,
                          'parameter-distributions_corn_IBO_EtOH_B.xlsx')
#: <scenario>.json records of each reproduction (reproduced / reloaded
#: values, recorded-row comparison), read by verify_opt_scenarios.py.
REPRODUCTION_DIR = os.path.join(_ANALYSES_DIR, 'results',
                                'opt_scenario_reproductions')
#: Non-kinetic decision entries reproduce_split12d_trial passes through in
#: result['applied']; the feeding strategy is read from result['feeding'].
FEEDING_KEYS = ('threshold_conc', 'target_delta', 'max_n_spikes')
#: Tracked metrics that describe the convergence path, not the point
#: (kinetic_optimization.REPRODUCTION_DIAGNOSTIC_METRICS); never compared.
DIAGNOSTIC_METRICS = ('spike_feed_residual', 'n_sims_run', 'final_drift')
_LOAD_RE = re.compile(r'_te\.(\w+)\s*=')

#%% Pure helpers (no model access; covered by the offline test)
def workbook_filename(scenario):
    """The scenario's workbook basename under WORKBOOK_DIR."""
    return f'parameter-distributions_corn_IBO_EtOH_{scenario}.xlsx'


def merge_applied(a_snapshot, applied):
    """The scenario's FULL kinetic state {name: float} over exactly
    `a_snapshot`'s names: the trial's `applied` values where sampled /
    grouped, the anchor snapshot elsewhere. FEEDING_KEYS entries of
    `applied` are dropped; any other name absent from the snapshot is a
    ValueError (a kinetic name the live model does not have)."""
    unknown = [name for name in applied
               if name not in a_snapshot and name not in FEEDING_KEYS]
    if unknown:
        raise ValueError(f'applied names {unknown} are neither kinetic '
                         'parameters of the anchor snapshot nor feeding keys '
                         f'{FEEDING_KEYS}')
    return {name: float(applied.get(name, value))
            for name, value in a_snapshot.items()}


def split_metric_warnings(metric_warnings, levels):
    """(blocking, tolerated): the reproduction's metric_warnings split by
    OBJECTIVE_REGISTRY level -- 'system' (TEA; hensmith-dependent) names are
    tolerated, everything else (kinetic-level registry metrics and the
    tracked-only fermentation metrics tau / n_glu_spikes / acetate) blocks.
    `levels` = {name: level} for the registry entries."""
    blocking = [name for name in metric_warnings
                if levels.get(name) != 'system']
    tolerated = [name for name in metric_warnings
                 if levels.get(name) == 'system']
    return blocking, tolerated


def write_workbook(template_path, out_path, full_applied, ibo_price,
                   provenance=None):
    """Copy `template_path` (the B workbook) to `out_path` with every kinetic
    row's Baseline <- full_applied[name] and a Triangular +/-20 %
    distribution, and the isobutanol-price row <- `ibo_price` with the
    template's relative spread preserved; every other row is kept. With
    `provenance`, each kinetic row's References cell becomes `provenance`
    (+ ' | template note: <old>' when the template had one), so a copied
    note such as B's "k_17 = 44 ..." cannot pass for the new Baseline's
    source. ValueError BEFORE anything is saved if a kinetic row's name is
    not in `full_applied`. Returns the number of kinetic rows written."""
    from openpyxl import load_workbook
    wb = load_workbook(template_path, data_only=True)
    ws = wb.active
    header = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value}
    col_load, col_base = header['Load statement'], header['Baseline']
    col_shape, col_refs = header['Shape'], header['References']
    col_low, col_mid, col_up = (header['Lower'], header['Midpoint'],
                                header['Upper'])
    kinetic_rows, price_rows = [], []
    for r in ws.iter_rows(min_row=2):
        load_stmt = r[col_load].value
        if not load_stmt:
            continue
        m = _LOAD_RE.search(str(load_stmt))
        if m:
            kinetic_rows.append((m.group(1), r))
        elif 'isobutanol_price' in str(load_stmt):
            price_rows.append(r)
    missing = [name for name, _ in kinetic_rows if name not in full_applied]
    if missing:
        raise ValueError(f'kinetic workbook rows {missing} are not in the '
                         'applied kinetic state; nothing written')
    for name, r in kinetic_rows:
        base = float(full_applied[name])
        r[col_base].value = base
        r[col_shape].value = 'Triangular'
        r[col_low].value = 0.8*base
        r[col_mid].value = 1.0*base
        r[col_up].value = 1.2*base
        if provenance:
            old = r[col_refs].value
            r[col_refs].value = (f'{provenance} | template note: {old}'
                                 if old else provenance)
    for r in price_rows:
        template_base = r[col_base].value
        r[col_base].value = float(ibo_price)
        if template_base:                  # preserve the relative spread
            scale = float(ibo_price)/float(template_base)
            for c in (col_low, col_mid, col_up):
                if r[c].value is not None:
                    r[c].value = float(r[c].value)*scale
    wb.save(out_path)
    return len(kinetic_rows)


def assert_reload_matches(reproduced, reloaded, tol):
    """{key: relative delta} over every key of `reproduced` (MPSPs
    'ethanol' / 'isobutanol' and the tracked metrics) between the
    reproduction and the reload-verify simulation (|a - b| / max(|a|, |b|),
    Python floats -- flexsolve's global np.seterr(invalid='raise') makes
    numpy-scalar NaN comparisons raise). A non-finite pair matches only if
    both are NaN or both the SAME infinity (delta NaN). RuntimeError naming
    every key whose delta exceeds `tol`, that is absent from `reloaded`, or
    whose finiteness differs."""
    deltas, bad = {}, []
    for key, a in reproduced.items():
        if key not in reloaded:
            bad.append(f'{key}: absent from the reload-verify simulation')
            continue
        a, b = float(a), float(reloaded[key])
        if not (math.isfinite(a) and math.isfinite(b)):
            deltas[key] = math.nan
            if not ((math.isnan(a) and math.isnan(b))
                    or (math.isinf(a) and a == b)):
                bad.append(f'{key}: reproduced {a!r} vs reloaded {b!r} '
                           '(non-finite mismatch)')
            continue
        scale = max(abs(a), abs(b))
        deltas[key] = 0.0 if scale == 0.0 else abs(a - b)/scale
        if deltas[key] > tol:
            bad.append(f'{key}: reproduced {a!r} vs reloaded {b!r} '
                       f'(rel {deltas[key]:.3g} > {tol:g})')
    if bad:
        raise RuntimeError('the written workbook does not reproduce the '
                           'trial: ' + '; '.join(bad))
    return deltas


def comparable_values(MPSPs, metrics):
    """The flat {key: float} assert_reload_matches compares: the two MPSPs
    ('ethanol', 'isobutanol') + every tracked metric but the convergence
    diagnostics."""
    values = dict(ethanol=float(MPSPs['ethanol']),
                  isobutanol=float(MPSPs['isobutanol']))
    values.update({name: float(value) for name, value in metrics.items()
                   if name not in DIAGNOSTIC_METRICS})
    return values


def format_scenario_spec(name, workbook, feeding, mpsps, objective_name,
                         objective_value):
    """The paste-ready scenarios.SCENARIOS entry. `feeding` is
    reproduce_split12d_trial's result['feeding'] (threshold / target /
    spike / max_n_spikes); the spike concentration is NOT written
    (spike_conc stays None = the load() default the preset pins)."""
    return (
        f"    '{name}': ScenarioSpec(\n"
        f"        name='{name}',\n"
        f"        workbook='{workbook}',\n"
        f"        max_n_spikes={int(feeding['max_n_spikes'])}, "
        f"threshold_conc={float(feeding['threshold'])!r}, "
        f"target_conc={float(feeding['target'])!r},\n"
        f"        expected={{'ethanol': {float(mpsps['ethanol'])!r}, "
        f"'isobutanol': {float(mpsps['isobutanol'])!r}}},\n"
        f"        objective_name={objective_name!r}, "
        f"objective_value={float(objective_value)!r}),")


#%% Simulation (needs the built model; called only from main)
def reload_verify(out_path, feeding):
    """Load the just-written workbook fresh (scenarios.load_scenario's
    recipe: parameters reset -> load_parameter_distributions ->
    metrics_at_baseline, which SETS the kinetics), set the trial's feeding,
    simulate under the still-active A-referenced burden and return
    comparable_values(MPSPs, every tracked metric)."""
    from biorefineries import isobutanol
    from biorefineries.isobutanol import kinetic_optimization as ko
    model = isobutanol.models.models_EtOH_IBO_corn.model
    handles = ko.get_handles()
    model.parameters = ()
    model.load_parameter_distributions(out_path,
                                       isobutanol.models.namespace_dict)
    model.metrics_at_baseline()
    handles['fbs_spec'].max_n_spikes = int(feeding['max_n_spikes'])
    handles['model_specification'](threshold_conc=feeding['threshold'],
                                   target_conc=feeding['target'],
                                   spike_conc=feeding['spike'])
    solution = handles['solve_TEA'](stream_IDs=('ethanol', 'isobutanol'),
                                    IRR_for_MPSP=0.15)
    handles['latest_TEA_solution'].update(solution)
    metrics = {name: getter(handles)
               for name, getter in ko.TRACKED_METRICS.items()}
    return comparable_values(solution['MPSPs'], metrics)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    parser.add_argument('scenario', choices=sorted(OPT_SCENARIOS))
    parser.add_argument('--study-name', default=None,
                        help='override the OPT_SCENARIOS campaign')
    parser.add_argument('--trial-number', type=int, default=None,
                        help='override the OPT_SCENARIOS trial')
    parser.add_argument('--anchor', default=ANCHOR_SCENARIO)
    args = parser.parse_args(argv)
    objective_name, study_name, trial_number = OPT_SCENARIOS[args.scenario]
    study_name = args.study_name or study_name
    trial_number = (trial_number if args.trial_number is None
                    else args.trial_number)
    workbook = workbook_filename(args.scenario)
    out_path = os.path.join(WORKBOOK_DIR, workbook)

    # 1. Build (default both-trains == IBO_EtOH-only at S201 split 1.0).
    from biorefineries import isobutanol
    isobutanol.load()
    from biorefineries.isobutanol import scenarios
    from biorefineries.isobutanol import kinetic_optimization as ko
    from biorefineries.isobutanol import system as ibo_system
    if objective_name not in ko.OBJECTIVE_REGISTRY:
        raise ValueError(f'{objective_name!r} is not in ko.OBJECTIVE_REGISTRY')

    # 2. Anchor snapshot: full live kinetics + isobutanol price.
    bundle = scenarios.load_scenario(args.anchor)
    a_snapshot = ko.discover_kinetic_parameters(ko.get_handles()['r_te'])
    ibo_price = float(bundle['model'].system.flowsheet.V514.isobutanol_price)
    print(f'Scenario-{args.anchor} snapshot: {len(a_snapshot)} kinetic '
          f'params; isobutanol price {ibo_price:.6g}')

    # 3. Reproduce the trial (read-only w.r.t. the campaign); require it clean.
    result = ko.reproduce_split12d_trial(
        args.anchor, study_name, trial_number, mode='both', burden=True,
        restore=True, metric_check_tol=METRIC_CHECK_TOL)
    levels = {name: entry['level']
              for name, entry in ko.OBJECTIVE_REGISTRY.items()}
    blocking, tolerated = split_metric_warnings(result['metric_warnings'],
                                                levels)
    problems = []
    if result['error'] is not None:
        problems.append(f"simulation error {result['error']}")
    if result['recorded_state'] != 'COMPLETE':
        problems.append(f"recorded state {result['recorded_state']!r}")
    if result['cross_check']['mismatches']:
        problems.append(f"cross-check mismatches "
                        f"{result['cross_check']['mismatches']} (the campaign "
                        'likely ran under another anchor / reference)')
    if blocking:
        problems.append(f'fermentation-level metric warnings {blocking}')
    if problems:
        raise RuntimeError('reproduction is not clean; nothing written: '
                           + '; '.join(problems))
    if tolerated:
        print(f'NOTE: system-level (TEA) metrics {tolerated} differ from the '
              f'recorded values by > {METRIC_CHECK_TOL:g} (hensmith-'
              'dependent; tolerated -- the pins are taken on the live model)')
    feeding = result['feeding']
    metrics = result['reproduced']['metrics']
    if objective_name not in metrics:
        raise KeyError(f'{objective_name!r} is not a tracked metric of the '
                       'reproduction')
    reproduced = comparable_values(result['reproduced']['MPSPs'], metrics)

    # 4. Full kinetic state: trial values over the anchor snapshot.
    full_applied = merge_applied(a_snapshot, result['applied'])

    # 5. Write the workbook (previous versions live in git).
    provenance = (f'{args.scenario}: trial {result["trial_number"]} of '
                  f'{os.path.basename(str(study_name))} (anchor scenario '
                  f'{args.anchor}; build_opt_split12d_workbooks.py)')
    n_kin = write_workbook(B_WORKBOOK, out_path, full_applied, ibo_price,
                           provenance=provenance)
    print(f'wrote {out_path} ({n_kin} kinetic rows)')

    # 6. Reload-verify under the still-active A-referenced burden.
    reloaded = reload_verify(out_path, feeding)
    deltas = assert_reload_matches(reproduced, reloaded, RELOAD_VERIFY_TOL)
    finite = [d for d in deltas.values() if not math.isnan(d)]
    print(f'reload-verify: {len(deltas)} values (MPSPs + tracked metrics), '
          f'max rel delta {max(finite, default=0.0):.2e} (tol '
          f'{RELOAD_VERIFY_TOL:g})')

    # 7. Record the reproduction (for the fresh-kernel scenario comparison)
    #    and print the registry entry.
    os.makedirs(REPRODUCTION_DIR, exist_ok=True)
    record_path = os.path.join(REPRODUCTION_DIR, f'{args.scenario}.json')
    with open(record_path, 'w') as f:
        json.dump(dict(scenario=args.scenario, study_name=study_name,
                       trial_number=result['trial_number'],
                       objective_name=objective_name, feeding=feeding,
                       reproduced=reproduced, reloaded=reloaded,
                       recorded_check=result['metric_check'],
                       tolerated_tea_warnings=tolerated),
                  f, indent=1)
    print(f'wrote {record_path}')
    print('\n# ---- paste into scenarios.SCENARIOS ----')
    print(format_scenario_spec(args.scenario, workbook, feeding,
                               reproduced, objective_name,
                               reproduced[objective_name]))
    ibo_system.set_active_burden(None)   # cleanliness -- nothing runs after
    return dict(result=result, full_applied=full_applied,
                reproduced=reproduced, reloaded=reloaded, deltas=deltas,
                out_path=out_path)


if __name__ == '__main__':
    main()
